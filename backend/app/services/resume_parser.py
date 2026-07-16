import io
import os
import re
import time
import zipfile
import logging
import requests
from typing import Tuple
import pdfplumber
from docx import Document
from app.config import settings

logger = logging.getLogger(__name__)

class ResumeParserError(Exception):
    pass

def validate_file(filename: str, file_size_bytes: int) -> None:
    """
    Validates file extension and size.
    Raises ResumeParserError if validation fails.
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext not in settings.allowed_extensions_list:
        raise ResumeParserError(
            f"Unsupported file format '{ext}'. Allowed formats: {settings.ALLOWED_RESUME_EXTENSIONS}"
        )
    
    max_bytes = settings.MAX_RESUME_FILE_SIZE_MB * 1024 * 1024
    if file_size_bytes > max_bytes:
        raise ResumeParserError(
            f"File size exceeds limit of {settings.MAX_RESUME_FILE_SIZE_MB}MB."
        )

def extract_text_from_pdf_plumber(file_bytes: bytes) -> str:
    """
    Extracts text from a PDF file using pdfplumber (fallback method).
    """
    text_content = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text()
                if text:
                    text_content.append(text)
    except Exception as e:
        raise ResumeParserError(f"Error parsing PDF file with pdfplumber: {str(e)}")
    
    if not text_content:
        raise ResumeParserError("No text content could be extracted from PDF.")
        
    return "\n".join(text_content)

def upload_to_tmpfiles(file_bytes: bytes) -> str:
    """
    Uploads file bytes to tmpfiles.org and extracts the direct download link.
    """
    try:
        # Step 1: Upload file to tmpfiles.org API
        url = "https://tmpfiles.org/api/v1/upload"
        files = {"file": ("resume.pdf", file_bytes, "application/pdf")}
        data = {"expire": 900}  # Expire in 15 minutes to preserve privacy
        
        response = requests.post(url, files=files, data=data, timeout=10)
        if response.status_code != 200:
            raise Exception(f"Upload failed with status code {response.status_code}: {response.text}")
            
        result = response.json()
        page_url = result.get("data", {}).get("url")
        if not page_url:
            raise Exception("No URL returned in upload response")
            
        # Step 2: Fetch the landing page HTML to extract the direct download link
        html_response = requests.get(page_url, headers={"User-Agent": "Mozilla/5.0"}, timeout=10)
        if html_response.status_code != 200:
            raise Exception(f"Failed to fetch landing page to parse direct link: {html_response.status_code}")
            
        html = html_response.text
        # Parse the direct download URL from the href tag
        match = re.search(r'href=\"(https://tmpfiles.org/dl/[^\"]+)\"', html)
        if not match:
            # Fall back to page URL with direct substitution if parser fails
            dl_url = page_url.replace("tmpfiles.org/", "tmpfiles.org/dl/")
            logger.warning(f"Regex match failed. Falling back to direct URL substitution: {dl_url}")
            return dl_url
            
        dl_url = match.group(1)
        logger.info(f"Direct download link prepared: {dl_url}")
        return dl_url
    except Exception as e:
        logger.error(f"Error uploading to tmpfiles.org: {str(e)}")
        raise

def parse_zip_for_markdown(zip_bytes: bytes) -> str:
    """
    Parses a downloaded zip archive in-memory to extract the main Markdown file.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
            # Look for any file ending with .md
            for filename in z.namelist():
                if filename.endswith(".md"):
                    with z.open(filename) as f:
                        return f.read().decode("utf-8")
        raise Exception("No Markdown file (.md) found inside the MinerU output zip")
    except Exception as e:
        logger.error(f"Error extracting markdown from ZIP: {str(e)}")
        raise

def extract_text_from_pdf_mineru(file_bytes: bytes) -> str:
    """
    Extracts text from a PDF file using the mineru.net Cloud API.
    Gracefully falls back to pdfplumber on failure.
    """
    token = settings.MINERU_API_TOKEN
    
    if settings.USE_MINERU_PARSER and token and token != "your_mineru_token_here" and not token.startswith("#"):
        try:
            logger.info("Initializing MinerU Cloud Parser...")
            
            # 1. Upload to tmpfiles.org to get a public direct link
            public_dl_url = upload_to_tmpfiles(file_bytes)
            
            # 2. Submit parsing task to mineru.net
            task_url = "https://mineru.net/api/v4/extract/task"
            headers = {
                "Authorization": f"Bearer {token}",
                "Content-Type": "application/json"
            }
            payload = {
                "url": public_dl_url,
                "model_version": "vlm",  # Recommended model version for high accuracy
                "is_ocr": True,
                "enable_table": True,
                "enable_formula": False  # Disable formulas for faster parsing of resumes
            }
            
            logger.info("Submitting task to MinerU cloud API...")
            response = requests.post(task_url, json=payload, headers=headers, timeout=15)
            if response.status_code != 200:
                raise Exception(f"Task submission failed: status code {response.status_code}, {response.text}")
                
            task_data = response.json()
            if task_data.get("code") != 0:
                raise Exception(f"MinerU returned error code {task_data.get('code')}: {task_data.get('msg')}")
                
            task_id = task_data.get("data", {}).get("task_id")
            if not task_id:
                raise Exception("No task_id returned by MinerU API")
                
            logger.info(f"Task submitted successfully. Task ID: {task_id}. Polling status...")
            
            # 3. Poll for status (max 20 attempts, checking every 2 seconds -> max 40 seconds)
            status_url = f"https://mineru.net/api/v4/extract/task/{task_id}"
            max_attempts = 20
            poll_interval = 2
            
            for attempt in range(max_attempts):
                time.sleep(poll_interval)
                status_response = requests.get(status_url, headers=headers, timeout=10)
                if status_response.status_code != 200:
                    logger.warning(f"Polling HTTP error {status_response.status_code}: {status_response.text}")
                    continue
                    
                status_data = status_response.json()
                if status_data.get("code") != 0:
                    raise Exception(f"MinerU polling error: {status_data.get('msg')}")
                    
                data = status_data.get("data", {})
                state = data.get("state")
                logger.info(f"MinerU polling attempt {attempt + 1}: state = {state}")
                
                if state in ["done", "success"]:
                    # Task completed successfully. Download zip file.
                    zip_url = data.get("full_zip_url")
                    if not zip_url:
                        raise Exception("Task completed but full_zip_url is missing")
                        
                    logger.info(f"Downloading parser output from ZIP URL...")
                    zip_res = requests.get(zip_url, timeout=20)
                    if zip_res.status_code != 200:
                        raise Exception(f"Failed to download zip file: status code {zip_res.status_code}")
                        
                    # Extract markdown
                    markdown_text = parse_zip_for_markdown(zip_res.content)
                    logger.info("Successfully extracted text from PDF via MinerU cloud API.")
                    return markdown_text
                    
                elif state == "failed":
                    raise Exception("MinerU task failed during cloud processing")
                    
            raise Exception("MinerU parsing timed out in cloud queue")
            
        except Exception as e:
            logger.error(f"MinerU Cloud Parser failed: {str(e)}. Falling back to pdfplumber.")
            
    else:
        logger.info("MinerU parser is disabled or no token is configured. Defaulting to pdfplumber.")
        
    return extract_text_from_pdf_plumber(file_bytes)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from a PDF file using a multi-tiered architecture:
    1. Tier 1: Try MarkItDown. Save output to markitdown_extracted_debug.md.
    2. If text length < 100 (scanned or flat PDF), fallback to Tier 2: MinerU Cloud API.
    3. If MinerU fails or is unconfigured, fallback to Tier 3: pdfplumber.
    """
    text = ""
    try:
        from markitdown import MarkItDown
        md = MarkItDown()
        result = md.convert_stream(io.BytesIO(file_bytes), file_extension=".pdf")
        text = result.text_content or ""
        logger.info(f"MarkItDown successfully processed PDF. Length: {len(text)}")
        
        # Save output to markitdown_extracted_debug.md as requested
        try:
            debug_path = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "markitdown_extracted_debug.md")
            with open(debug_path, "w", encoding="utf-8") as f:
                f.write(text)
            logger.info(f"Saved MarkItDown extraction text to debug file: {debug_path}")
        except Exception as write_err:
            logger.warning(f"Failed to write MarkItDown output to debug file: {str(write_err)}")
            
    except Exception as e:
        logger.error(f"MarkItDown native conversion failed: {str(e)}")
        
    # Check if we should fall back to MinerU / pdfplumber
    if len(text.strip()) < 100:
        logger.warning("Flat file or scan detected (or MarkItDown failed). Triggering MinerU cloud parser fallback...")
        return extract_text_from_pdf_mineru(file_bytes)
        
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from a DOCX file using python-docx.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs]
        
        # Also extract table text to avoid missing contents in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
                    
        text = "\n".join(p for p in paragraphs if p.strip())
        if not text:
            raise ResumeParserError("No text content could be extracted from DOCX.")
        return text
    except Exception as e:
        if not isinstance(e, ResumeParserError):
            raise ResumeParserError(f"Error parsing DOCX file: {str(e)}")
        raise

def parse_resume(filename: str, file_bytes: bytes) -> str:
    """
    Validates and extracts raw text from PDF or DOCX file bytes.
    """
    validate_file(filename, len(file_bytes))
    
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        text = extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        text = extract_text_from_docx(file_bytes)
    else:
        raise ResumeParserError(f"Unsupported file format: {ext}")
        
    # Clean up PDF glyph artifacts like (cid:263)
    import re
    text = re.sub(r'\(cid:\d+\)', ' ', text)
    # Collapse multiple consecutive spaces
    text = re.sub(r'[ \t]+', ' ', text)
    
    return text
